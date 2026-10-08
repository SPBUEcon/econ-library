# -*- coding: utf-8 -*-
"""Собирает представления тезисов из md-источников.

Источник правды — карточки theses/th*/README.md и сборки theses/assemblies/*.md.
Всё остальное — представления в едином стиле (theme.css). У тезиса и у сборки их два:

- слайды (thNN-…/slides.html, assemblies/<id>-slides.html) — чистый контент: показ, PDF, PPTX;
- карточка (thNN-…/index.html, assemblies/<id>.html) — слайды с подсветкой проверок, комментарии,
  вопросы, задания, связи.

Служебная строка слайда — индекс тезиса (ссылка на карточку) и дата обновления.

    python theses/site/build.py            # HTML-страницы, витрина, CSV-индексы

    python theses/site/build.py            # HTML-страницы, витрина, CSV-индексы
    python theses/site/build.py --pdf      # + PDF (слайды 16:9, по странице на слайд)
    python theses/site/build.py --pptx     # + PPTX (кадры слайдов + заметки докладчика)
    python theses/site/build.py --export   # всё сразу

HTML-страницы самодостаточны (стиль и скрипт встроены), работают по file:// и на GitHub Pages.
PDF и PPTX делаются из тех же HTML через headless Chrome / Edge и кладутся в theses/_export/
(в git не попадает). Путь к браузеру можно задать переменной THESES_BROWSER.
"""
import csv
import glob
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent                     # theses/
REPO = ROOT.parent
EXPORT = ROOT / "_export"
GH = "https://github.com/SPBUEcon/econ-library/"
PAGES = "https://spbuecon.github.io/econ-library/"   # ссылки в PDF и PPTX ведут сюда
FOOT = "Мастерская «Технологии в экономике» · ЭФ СПбГУ"
FONTS = ('<link rel="preconnect" href="https://fonts.googleapis.com">\n'
         '<link href="https://fonts.googleapis.com/css2?family=Unbounded:wght@500;700'
         '&family=IBM+Plex+Sans:ital,wght@0,400;0,500;0,600;1,500&display=swap" rel="stylesheet">')

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

WARN = []


def warn(msg):
    WARN.append(msg)


def split(v):
    return [x.strip() for x in (v or "").split(";") if x.strip()]


def esc(s):
    return html.escape(s, quote=True)


# ---------------------------------------------------------------- md → html

class Ctx:
    """Где лежит md-источник и где окажется html (нужно для перевода ссылок)."""

    def __init__(self, src_dir, out_dir, theses=None, clean=False):
        self.src_dir = Path(src_dir)
        self.out_dir = Path(out_dir)
        self.theses = theses or {}     # путь README.md → путь index.html и т.п.
        self.clean = clean             # чистые слайды: без подсветки проверок и пометок


def link_target(url, ctx):
    """Относительная ссылка из md → ссылка из html.

    md тезисов и сборок → их html; прочие .md/.csv и папки → GitHub; остальное — относительно.
    """
    if re.match(r"^[a-z]+:|^#|^mailto:", url):
        return url
    path, _, anchor = url.partition("#")
    if not path:
        return url
    target = (ctx.src_dir / path).resolve()
    if target in ctx.theses:
        out = ctx.theses[target]
    elif target.suffix in (".md", ".csv") or target.is_dir() or not target.suffix:
        rel = target.relative_to(REPO).as_posix() if target.is_relative_to(REPO) else path
        kind = "tree" if target.is_dir() else "blob"
        return f"{GH}{kind}/main/{rel}" + (f"#{anchor}" if anchor else "")
    else:
        out = target
    rel = os.path.relpath(out, ctx.out_dir).replace(os.sep, "/")
    return rel + (f"#{anchor}" if anchor else "")


CHK = re.compile(r"^\(?(со слайда\s*[—-]\s*)?проверить\)?$")


def inline(s, ctx):
    keep = []

    def stash(h):
        keep.append(h)
        return f"\x00{len(keep) - 1}\x00"

    s = re.sub(r"!\[([^\]]*)\]\(([^)\s]+)\)",
               lambda m: stash(f'<img src="{esc(link_target(m.group(2), ctx))}" alt="{esc(m.group(1))}">'), s)
    s = re.sub(r"\[((?:[^\[\]]|\[[^\]]*\])*)\]\(([^)\s]+)\)",
               lambda m: stash(f'<a href="{esc(link_target(m.group(2), ctx))}">{inline(m.group(1), ctx)}</a>'), s)
    s = re.sub(r"`([^`]+)`", lambda m: stash(f"<code>{esc(m.group(1))}</code>"), s)
    # <mark>…</mark> — утверждение, которое нужно проверить: в карточке подсвечено, в слайдах — обычный текст
    s = s.replace("<mark>", "\x01").replace("</mark>", "\x02")
    s = esc(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)

    def em(m):
        t = m.group(1)
        if CHK.match(t):
            return "" if ctx.clean else '<span class="chk">проверить</span>'
        return f"<em>{t}</em>"

    s = re.sub(r"(?<![*\w])\*(?![\s*])(.+?)(?<![\s*])\*(?![*\w])", em, s)
    s = s.replace("\x01", "" if ctx.clean else "<mark>").replace("\x02", "" if ctx.clean else "</mark>")
    return re.sub(r"\x00(\d+)\x00", lambda m: keep[int(m.group(1))], s)


def slug(text):
    """Якорь заголовка, как у GitHub: нижний регистр, пробелы → дефис, без пунктуации."""
    t = re.sub(r"[`*]|\[|\]\([^)]*\)", "", text).strip().lower()
    return re.sub(r"[^\w\- ]", "", t).replace(" ", "-")


LIST = re.compile(r"^\s*([-*]|\d+[.)])\s+(.*)$")


def blocks(md):
    """md → список блоков (kind, data). Подмножество md, которого хватает карточкам."""
    lines = md.splitlines()
    out, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
        elif ln.strip().startswith("<!--"):
            while "-->" not in lines[i]:
                i += 1
            i += 1
        elif ln.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            lang = ln[3:].strip()
            out.append(("cycle" if lang == "cycle" else "pre", "\n".join(lines[i + 1:j])))
            i = j + 1
        elif re.match(r"^#{1,6} ", ln):
            lvl = len(ln) - len(ln.lstrip("#"))
            out.append(("h", (lvl, ln[lvl:].strip())))
            i += 1
        elif ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(("table", (rows[0], rows[2:])))
        elif ln.startswith(">"):
            inner = []
            while i < len(lines) and lines[i].startswith(">"):
                inner.append(re.sub(r"^> ?", "", lines[i]))
                i += 1
            out.append(("quote", "\n".join(inner)))
        elif LIST.match(ln):
            ordered = LIST.match(ln).group(1)[0].isdigit()
            items = []
            while i < len(lines) and lines[i].strip():
                m = LIST.match(lines[i])
                if m:
                    items.append(m.group(2))
                elif lines[i].startswith((" ", "\t")) and items:
                    items[-1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            out.append(("ol" if ordered else "ul", items))
        else:
            para = []
            while (i < len(lines) and lines[i].strip() and not lines[i].startswith(("|", ">", "```", "#", "<!--"))
                   and not LIST.match(lines[i])):
                para.append(lines[i].strip())
                i += 1
            out.append(("p", " ".join(para)))
    return out


def render(bl, ctx, layout=""):
    h = []
    for kind, d in bl:
        if kind == "p":
            h.append(f"<p>{inline(d, ctx)}</p>")
        elif kind == "h":
            lvl = min(d[0] + 1, 6)
            h.append(f'<h{lvl} id="{slug(d[1])}">{inline(d[1], ctx)}</h{lvl}>')
        elif kind == "pre":
            h.append(f"<pre>{esc(d)}</pre>")
        elif kind == "cycle":
            h.append(cycle_html(d, ctx))
        elif kind == "quote":
            h.append(f"<blockquote>{render(blocks(d), ctx)}</blockquote>")
        elif kind in ("ul", "ol"):
            h.append(f"<{kind}>" + "".join(f"<li>{inline(x, ctx)}</li>" for x in d) + f"</{kind}>")
        elif kind == "table" and layout == "numbers" and all(len(r) > 1 and number(r[1]) is not None for r in d[1]):
            # таблица с числами во втором столбце — диаграмма; таблицы подряд — диаграммы рядом
            if h and isinstance(h[-1], list):
                h[-1].append(d)
            else:
                h.append([d])
        elif kind == "table" and layout in ("columns", "timeline"):
            head, rows = d
            cols = []
            for k, name in enumerate(head):
                title, _, sub = name.partition(" · ")
                items = [r[k] for r in rows if k < len(r) and r[k]]
                cols.append(f'<div class="col"><div class="col-h"><b>{inline(title, ctx)}</b>'
                            + (f"<span>{inline(sub, ctx)}</span>" if sub else "") + "</div><ul>"
                            + "".join(f"<li>{inline(x, ctx)}</li>" for x in items) + "</ul></div>")
            h.append('<div class="cols">' + "".join(cols) + "</div>")
        elif kind == "table":
            head, rows = d
            h.append("<table><thead><tr>" + "".join(f"<th>{inline(c, ctx)}</th>" for c in head) + "</tr></thead><tbody>"
                     + "".join("<tr>" + "".join(f"<td>{inline(c, ctx)}</td>" for c in r) + "</tr>" for r in rows)
                     + "</tbody></table>")
    return "\n".join(charts_html(x, ctx) if isinstance(x, list) else x for x in h)


def number(s):
    """«4,03 трлн», «<mark>211,0</mark>» → 4.03, 211.0; без числа — None."""
    m = re.search(r"\d[\d\s]*(?:[.,]\d+)?", re.sub(r"<[^>]+>", "", s))
    return float(re.sub(r"\s", "", m.group(0)).replace(",", ".")) if m else None


def charts_html(tables, ctx):
    """Диаграммы рядом. У всех одна единица (второй столбец шапки) — общий масштаб, чтобы длины сравнивались."""
    top = None
    if len({t[0][1] if len(t[0]) > 1 else "" for t in tables}) == 1:
        top = max(number(r[1]) or 0 for t in tables for r in t[1]) or None
    return '<div class="charts">' + "".join(bars_html(*t, ctx, top) for t in tables) + "</div>"


def bars_html(head, rows, ctx, top=None):
    """Таблица «название | число | пометка» → диаграмма-«леденец»: тонкая линия и точка на конце.

    Шапка — подпись диаграммы: «Выручка за 2025 год | млрд ₽». Длина линии — доля от максимума (top).
    """
    vals = [number(r[1]) if len(r) > 1 else None for r in rows]
    top = top or max([v for v in vals if v] or [1])
    cap = " · ".join(inline(c, ctx) for c in head[:2] if c)
    out = []
    for r, v in zip(rows, vals):
        w = (v or 0) / top * 78                # место справа — под значение
        note = f'<i>{inline(r[2], ctx)}</i>' if len(r) > 2 and r[2] else ""
        out.append(f'<div class="b-row"><span class="b-lab">{inline(r[0], ctx)}</span><span class="b-track">'
                   f'<span class="b-line" style="width:{w:.1f}%"><b class="b-val">'
                   f'{inline(r[1] if len(r) > 1 else "", ctx)}{note}</b></span></span></div>')
    return f'<div class="bars"><div class="b-cap">{cap}</div>{"".join(out)}</div>'


CYCLE_STEP = re.compile(r"\s*-([^->]*)->\s*")


def cycle_node(n, ctx, extra=""):
    name, _, sub = n.partition(" · ")      # «Логистика · передача» → узел и подпись под ним
    return (f'<span class="cy-node{extra}">{inline(name, ctx)}'
            + (f'<small>{inline(sub, ctx)}</small>' if sub else "") + "</span>")


def cycle_arrow(verb, ctx):
    label = inline(verb, ctx) if verb.strip() else ""
    return '<span class="cy-arr">' + (f"<i>{label}</i>" if label else "") + "</span>"


def cycle_html(src, ctx):
    """```cycle: «A -глагол-> B -глагол-> C -глагол-> A» → схема: узлы, стрелки с подписями, обратная дуга.

    Первая строка — цикл (или цепочка, если последний узел не повторяет первый). Следующие строки вида
    «C -глагол-> D», где C — последний узел цикла, — ответвление: узел D дорисовывается справа от цикла.
    На GitHub блок читается как текст, в HTML — как схема на тонких линиях.
    """
    lines = [ln.strip() for ln in src.splitlines() if ln.strip()]
    parts = CYCLE_STEP.split(lines[0])
    nodes, verbs = parts[0::2], parts[1::2]
    loop = len(nodes) > 2 and nodes[-1].partition(" · ")[0] == nodes[0].partition(" · ")[0]
    back = verbs.pop() if loop else ""
    if loop:
        nodes.pop()
    row = []
    for k, n in enumerate(nodes):
        row.append(cycle_node(n, ctx))
        if k < len(verbs):
            row.append(cycle_arrow(verbs[k], ctx))
    ext = 0
    for ln in lines[1:]:
        bp = CYCLE_STEP.split(ln)
        if len(bp) == 3 and bp[0].partition(" · ")[0] == nodes[-1].partition(" · ")[0]:
            row += [cycle_arrow(bp[1], ctx), cycle_node(bp[2], ctx, " cy-ext")]
            ext += 1
        else:
            warn(f"cycle: ответвление «{ln}» должно начинаться с последнего узла цикла «{nodes[-1]}»")
    cls = ("cycle loop" if loop else "cycle chain") + (f" ext{ext}" if ext else "")
    return (f'<div class="{cls}"><div class="cy-row">' + "".join(row) + "</div>"
            + (f'<div class="cy-back"><i>{inline(back, ctx)}</i></div>' if loop else "") + "</div>")


def md_html(md, ctx, layout="", skip_quotes=False):
    bl = blocks(md)
    if skip_quotes:
        bl = [b for b in bl if b[0] != "quote"]
    return render(bl, ctx, layout)


def plain(md):
    """md → простой текст (заметки PPTX)."""
    t = re.sub(r"<!--.*?-->|</?mark>", "", md, flags=re.S)
    t = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", t)
    t = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", t)
    t = re.sub(r"[*`]|^> ?|^#+ ", "", t, flags=re.M)
    t = re.sub(r"^\|[\s:|-]+\|$", "", t, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", t).strip()


# ---------------------------------------------------------------- источники

def front(text):
    meta = {}
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if m:
        for ln in m.group(1).splitlines():
            k, _, v = ln.partition(":")
            if k.strip():
                meta[k.strip()] = v.strip()
        text = text[m.end():]
    return meta, text


def sections(body):
    """→ (h1, lead_md, [(заголовок, md)])"""
    h1 = re.search(r"^# (.+)$", body, re.M)
    rest = body[h1.end():] if h1 else body
    parts = re.split(r"^## (.+)$", rest, flags=re.M)
    return (h1.group(1).strip() if h1 else ""), parts[0], list(zip(parts[1::2], parts[2::2]))


SLIDE = re.compile(r"^Слайд\s*\d+\s*[.:·]?\s*(.*)$")


def load_thesis(readme):
    text = Path(readme).read_text(encoding="utf-8").replace("\r\n", "\n")
    meta, body = front(text)
    h1, lead, secs = sections(body)
    t = {"meta": meta, "h1": h1, "lead": lead, "slides": [], "sections": [],
         "dir": Path(readme).parent, "src": Path(readme).resolve()}
    for title, md in secs:
        m = SLIDE.match(title)
        if not m:
            t["sections"].append((title, md))
            continue
        def directive(name):
            d = re.search(r"<!--\s*" + name + r":\s*(\S+)\s*-->", md)
            return d.group(1) if d else ""

        lay = directive("layout")
        content, notes = (re.split(r"<!--\s*notes\s*-->", md, maxsplit=1) + [""])[:2]
        t["slides"].append({"title": m.group(1).strip(), "layout": lay,
                            "logo": directive("logo"), "bg": directive("bg"),
                            "md": content, "notes": notes.strip()})
    tid = meta.get("id", "")
    if not tid:
        warn(f"{readme}: нет id во frontmatter")
    if not meta.get("formula"):
        warn(f"{tid}: нет формулы — тезис должен укладываться в одну фразу")
    if not 1 <= len(t["slides"]) <= 3:
        warn(f"{tid}: слайдов {len(t['slides'])}, а тезис — это 1–3 слайда")
    t["id"], t["title"] = tid, meta.get("title", h1)
    t["out"] = t["dir"] / "index.html"
    t["deck"] = t["dir"] / "slides.html"
    return t


ITEM = re.compile(r"`(th\d+)`")


def load_assembly(path):
    """Сборка: разделы со списками `thNN` — группы тезисов, остальные разделы — текст.

    Сборки лежат в theses/assemblies/<id>.md (→ <id>.html, <id>-slides.html) или, для кейса, в
    README.md папки кейса с `kind: кейс` (→ index.html, slides.html рядом).
    """
    text = Path(path).read_text(encoding="utf-8").replace("\r\n", "\n")
    meta, body = front(text)
    h1, lead, secs = sections(body)
    parts, groups = [], []
    for title, md in secs:
        items = []
        for kind, d in blocks(md):
            if kind in ("ul", "ol"):
                for x in d:
                    m = ITEM.search(x)
                    items.append({"id": m.group(1) if m else "", "text": ITEM.sub("", x, count=1).strip() if m else x})
        if any(i["id"] for i in items):
            groups.append((title, items))
            parts.append(("group", title, items))
        else:
            parts.append(("text", title, md))
    p = Path(path)
    readme = p.name == "README.md"
    return {"meta": meta, "h1": h1, "lead": lead, "groups": groups, "parts": parts,
            "id": meta.get("id", p.parent.name if readme else p.stem),
            "title": meta.get("title", h1), "src": p.resolve(), "dir": p.parent,
            "out": p.parent / "index.html" if readme else p.with_suffix(".html"),
            "deck": p.parent / "slides.html" if readme else p.with_name(p.stem + "-slides.html")}


def case_sources():
    """README кейсов дисциплин, размеченные как сборка (`kind: кейс` во frontmatter)."""
    out = []
    for p in sorted(glob.glob(str(REPO / "disciplines" / "*" / "cases" / "*" / "README.md"))):
        meta, _ = front(Path(p).read_text(encoding="utf-8").replace("\r\n", "\n"))
        if meta.get("kind") == "кейс":
            out.append(p)
    return out


# ---------------------------------------------------------------- страницы

DECK_SCRIPT = """<script>
(function(){
  var slides=[].slice.call(document.querySelectorAll('main .slide')), stage=document.querySelector('.stage'), i=0;
  var links=[].slice.call(document.querySelectorAll('a[data-pages]')), q=new URLSearchParams(location.search);
  function web(on){links.forEach(function(a){if(!a.dataset.local)a.dataset.local=a.getAttribute('href');
    a.setAttribute('href',on?a.dataset.pages:a.dataset.local);});}
  function show(k){i=Math.max(0,Math.min(slides.length-1,k));stage.innerHTML='';stage.appendChild(slides[i].cloneNode(true));}
  if(q.has('export'))web(true);
  if(q.has('frame')){document.body.classList.add('frame');show((+q.get('frame')||1)-1);return;}
  window.addEventListener('beforeprint',function(){web(true);});
  window.addEventListener('afterprint',function(){if(!q.has('export'))web(false);});
  if(q.has('print'))(document.fonts?document.fonts.ready:Promise.resolve()).then(function(){setTimeout(function(){window.print();},300);});
  function start(k){document.body.classList.add('present');show(k||0);
    if(document.documentElement.requestFullscreen)document.documentElement.requestFullscreen().catch(function(){});}
  function stop(){document.body.classList.remove('present');if(document.fullscreenElement)document.exitFullscreen();}
  if(q.has('present'))start((+q.get('present')||1)-1);
  function current(){var k=0;slides.forEach(function(s,j){if(s.getBoundingClientRect().top<innerHeight/2)k=j;});return k;}
  [].forEach.call(document.querySelectorAll('[data-present]'),function(b){b.onclick=function(){start(current());};});
  [].forEach.call(document.querySelectorAll('[data-print]'),function(b){b.onclick=function(){window.print();};});
  slides.forEach(function(s,k){s.addEventListener('dblclick',function(){start(k);});});
  stage.addEventListener('click',function(e){if(e.target.closest('a'))return;show(e.clientX>innerWidth/3?i+1:i-1);});
  document.addEventListener('keydown',function(e){
    if(!document.body.classList.contains('present')){
      if(['f','F','а','А'].indexOf(e.key)>=0&&!e.ctrlKey&&!e.metaKey)start(current());return;}
    if(['ArrowRight','ArrowDown','PageDown',' ','Enter'].indexOf(e.key)>=0){show(i+1);e.preventDefault();}
    else if(['ArrowLeft','ArrowUp','PageUp','Backspace'].indexOf(e.key)>=0){show(i-1);e.preventDefault();}
    else if(['f','F','а','А'].indexOf(e.key)>=0&&!document.fullscreenElement)document.documentElement.requestFullscreen().catch(function(){});
    else if(e.key==='Escape')stop();});
  document.addEventListener('fullscreenchange',function(){if(!document.fullscreenElement)document.body.classList.remove('present');});
})();
</script>"""


def head(title, extra_css=""):
    css = (HERE / "theme.css").read_text(encoding="utf-8")
    return f"""<!DOCTYPE html>
<!-- Сгенерировано theses/site/build.py из md-источников. Не править руками: правьте md и пересоберите. -->
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{esc(title)}</title>
{FONTS}
<style>
{css}{extra_css}</style></head>"""


def page(title, hero, main):
    """Карточка: шапка + текст."""
    return head(title) + f"""
<body>
<header class="hero"><div class="wrap">
{hero}
</div></header>
<main><div class="wrap">
{main}
</div></main>
{CARD_SCRIPT}
</body></html>
"""


def deck_page(title, slides):
    """Чистые слайды: только слайды; показ, печать 16:9."""
    return head(title, "@page{size:1280px 720px;margin:0}\n") + f"""
<body class="deckpage">
<main><div class="deck">
{slides}
</div></main>
<div class="fab"><button data-present title="Показ: F или двойной клик по слайду">▶</button><button data-print title="Печать / PDF">⎙</button></div>
<div class="stage"></div>
{DECK_SCRIPT}
</body></html>
"""


def rel(target, from_dir):
    return os.path.relpath(target, from_dir).replace(os.sep, "/")


def pages_url(path):
    return PAGES + Path(path).relative_to(REPO).as_posix()


def tools(deck, base, src):
    href = rel(deck, base)
    return (f'<div class="tools"><a class="btn" href="{href}">▶ Слайды</a>'
            f'<a class="btn ghost" href="{href}?print">⎙ PDF</a>'
            f'<a class="btn ghost" href="{GH}blob/main/{src.relative_to(REPO).as_posix()}">md-источник</a></div>')


def s_top(label, link, ctx, updated):
    return (f'<div class="s-top"><a class="s-id" href="{rel(link, ctx.out_dir)}" data-pages="{pages_url(link)}"'
            f' title="Открыть карточку тезиса">{esc(label)}</a><span class="s-upd">обновлено {esc(updated)}</span></div>')


def slide_html(t, k, ctx, link, num, total):
    """Слайд. Служебная строка — индекс тезиса (ссылка link) и дата обновления, внизу — номер."""
    s = t["slides"][k]
    lay = s["layout"]
    bg = f' style="background-image:url(\'{esc(link_target(s["bg"], ctx))}\')"' if s["bg"] else ""
    logo = f'<img class="s-logo" src="{esc(link_target(s["logo"], ctx))}" alt="">' if s["logo"] else ""
    return (f'<section class="slide{" l-" + lay if lay else ""}"{bg}><div class="s-in">'
            + s_top(f'{t["id"]} · {t["title"]}', link, ctx, t["meta"].get("updated", ""))
            + f'<div class="s-head"><h2 class="s-title">{inline(s["title"], ctx)}</h2>{logo}</div>'
            f'<div class="s-body">{md_html(s["md"], ctx, lay)}</div>'
            f'<div class="s-foot"><span>{FOOT}</span><span>{num} / {total}</span></div>'
            "</div></section>")


def slides_seq(t, ctx, link, notes=False, start=0, total=None):
    h = []
    for k, s in enumerate(t["slides"]):
        h.append(slide_html(t, k, ctx, link, start + k + 1, total or len(t["slides"])))
        if notes and s["notes"]:
            h.append(f'<details class="notes doc"><summary>Комментарий к слайду {k + 1}</summary>'
                     f'{md_html(s["notes"], ctx)}</details>')
    return "\n".join(h)


def carousel(t, ctx, link, notes=False):
    """Компактная лента для карточки: один слайд на виду, ‹ ›, список слайдов, ▶ — показ с текущего слайда."""
    n = len(t["slides"])
    play = rel(t["deck"], ctx.out_dir)
    track = "".join(f'<div class="cr-item{" on" if k == 0 else ""}">{slide_html(t, k, ctx, link, k + 1, n)}</div>'
                    for k in range(n))
    nav = ('<button class="cr-btn cr-prev" aria-label="Предыдущий слайд">‹</button>'
           '<button class="cr-btn cr-next" aria-label="Следующий слайд">›</button>') if n > 1 else ""
    items = "".join(f'<li class="{"on" if k == 0 else ""}"><b>{k + 1}</b>{inline(s["title"], ctx)}</li>'
                    for k, s in enumerate(t["slides"]))
    com = ""
    if notes:
        com = "".join(f'<div class="cr-note doc{" on" if k == 0 else ""}">'
                      + (md_html(s["notes"], ctx) if s["notes"] else '<p class="empty">Комментария нет.</p>') + "</div>"
                      for k, s in enumerate(t["slides"]))
        com = f'<div class="cr-notes"><div class="cr-cap">Комментарий к слайду</div>{com}</div>'
    return (f'<div class="carousel" data-n="{n}"><div class="cr-stage">{track}{nav}</div>'
            f'<aside class="cr-side"><a class="btn cr-play" href="{play}?present=1" data-play="{play}">▶ Показ</a>'
            f'<ol class="cr-list">{items}</ol>{com}</aside></div>')


CARD_SCRIPT = """<script>
[].forEach.call(document.querySelectorAll('.carousel'),function(c){
  var items=c.querySelectorAll('.cr-item'),li=c.querySelectorAll('.cr-list li'),notes=c.querySelectorAll('.cr-note'),
      play=c.querySelector('.cr-play'),i=0;
  function go(k){i=(k+items.length)%items.length;
    [items,li,notes].forEach(function(l){[].forEach.call(l,function(e,j){e.classList.toggle('on',j===i);});});
    play.setAttribute('href',play.dataset.play+'?present='+(i+1));}
  var p=c.querySelector('.cr-prev'),n=c.querySelector('.cr-next');
  if(p)p.onclick=function(){go(i-1);}; if(n)n.onclick=function(){go(i+1);};
  [].forEach.call(li,function(e,j){e.onclick=function(){go(j);};});
  c.tabIndex=0;
  c.addEventListener('keydown',function(e){if(e.key==='ArrowRight'){go(i+1);e.preventDefault();}
    else if(e.key==='ArrowLeft'){go(i-1);e.preventDefault();}});
  [].forEach.call(items,function(e){e.addEventListener('dblclick',function(){location.href=play.getAttribute('href');});});
  c.go=go;
});
// старые якоря (#formula и т.п.) ведут на слайд тезиса в карусели
function hashSlide(){var a=location.hash&&document.getElementById(decodeURIComponent(location.hash.slice(1)));
  if(!a||!a.dataset.slide)return;var c=a.parentNode.querySelector('.carousel');if(c&&c.go){c.go(+a.dataset.slide-1);c.scrollIntoView();}}
window.addEventListener('hashchange',hashSlide);hashSlide();
</script>"""


def tags(meta, keys):
    out = []
    for k in keys:
        for v in split(meta.get(k)):
            out.append(f'<span class="tag">{esc(v)}</span>')
    return '<div class="tags">' + "".join(out) + "</div>"


def thesis_page(t, links, member_of):
    """Карточка тезиса: слайды с подсветкой проверок, комментарии, вопросы, задания, связи."""
    ctx = Ctx(t["dir"], t["dir"], links)
    m = t["meta"]
    hero = (f'<div class="crumbs"><a href="../site/index.html">Тезисы</a> · {esc(t["id"])}'
            + "".join(f' · <a href="{rel(a["out"], t["dir"])}">{esc(a["title"])}</a>' for a in member_of) + "</div>"
            f'<div class="eyebrow">Карточка тезиса · {esc(m.get("kind", ""))} · {len(t["slides"])} сл.</div>'
            f'<h1>{esc(t["title"])}</h1><div class="formula">{inline(m.get("formula", ""), ctx)}</div>'
            + tags(m, ["blocks", "refs"])
            + f'<div class="tags"><span class="tag">{esc(m.get("status", ""))} · {esc(m.get("version", ""))}'
              f' · обновлено {esc(m.get("updated", ""))}</span></div>' + tools(t["deck"], t["dir"], t["src"]))
    legend = ""
    if any("<mark>" in s["md"] for s in t["slides"]):
        legend = ('<p class="legend"><mark>Подсвечено</mark> то, что нужно проверить: список — в разделе '
                  '<a href="#что-проверить">«Что проверить»</a>. В <a href="slides.html">слайдах</a> '
                  'подсветки нет.</p>')
    main = (f'<div class="lead doc">{md_html(t["lead"], ctx, skip_quotes=True)}</div>{legend}'
            + carousel(t, ctx, t["deck"], notes=True)
            + '<div class="doc">' + "".join(f'<h2 id="{slug(title)}">{inline(title, ctx)}</h2>{md_html(md, ctx)}'
                                            for title, md in t["sections"]) + "</div>")
    return page(f'{t["id"]} · {t["title"]}', hero, main)


def thesis_slides(t, links):
    """Чистые слайды тезиса: индекс на слайде ведёт в карточку."""
    ctx = Ctx(t["dir"], t["dir"], links, clean=True)
    return deck_page(f'{t["id"]} · {t["title"]}', slides_seq(t, ctx, t["out"]))


def members(a, theses):
    return [theses[i["id"]] for _, its in a["groups"] for i in its if i["id"] in theses]


def anchors(meta):
    """`anchors: formula=th01:2; numbers=цифры-и-статус-проверки` → {тезис: [(якорь, слайд)]}, {раздел: [якорь]}.

    Сохраняет старые ссылки на страницу: якорь на слайд тезиса листает карусель, якорь на раздел — синоним.
    """
    to_slide, to_sec = {}, {}
    for pair in split(meta.get("anchors")):
        name, _, target = pair.partition("=")
        tid, _, num = target.strip().partition(":")
        if num:
            to_slide.setdefault(tid, []).append((name.strip(), int(num)))
        else:
            to_sec.setdefault(target.strip(), []).append(name.strip())
    return to_slide, to_sec


def assembly_page(a, theses, links):
    """Карточка сборки: разделы по порядку; группа — оглавление и карусели тезисов, текст — как есть."""
    ctx = Ctx(a["dir"], a["dir"], links)
    m = a["meta"]
    to_slide, to_sec = anchors(m)
    body = []
    for kind, title, data in a["parts"]:
        sid = slug(title)
        alias = "".join(f'<span class="anchor" id="{esc(x)}"></span>' for x in to_sec.get(sid, []))
        head = f'{alias}<h2 id="{sid}">{inline(title, ctx)}</h2>'
        if kind == "text":
            body.append(f'<div class="doc">{head}{md_html(data, ctx)}</div>')
            continue
        lis, units = [], []
        for it in data:
            t = theses.get(it["id"])
            if t:
                href = rel(t["out"], a["dir"])
                lis.append(f'<li><a href="{href}"><code>{t["id"]}</code> {esc(t["title"])}</a> — '
                           f'{inline(t["meta"].get("formula", ""), ctx)}</li>')
                tctx = Ctx(t["dir"], a["dir"], links, clean=True)
                marks = "".join(f'<span class="anchor" id="{esc(x)}" data-slide="{n}"></span>'
                                for x, n in to_slide.get(t["id"], []))
                units.append(f'<div class="unit" id="{t["id"]}">{marks}<div class="unit-h"><h3><a href="{href}">'
                             f'{t["id"]} · {esc(t["title"])}</a></h3><span class="fx">'
                             f'{inline(t["meta"].get("formula", ""), tctx)}</span></div>'
                             f'{carousel(t, tctx, t["out"])}</div>')
            else:
                if it["id"]:
                    warn(f'{a["id"]}: тезис {it["id"]} не найден')
                lis.append(f'<li class="planned">{inline(it["text"], ctx)} <span class="chk">планируется</span></li>')
        body.append(f'<div class="doc">{head}<ol>{"".join(lis)}</ol></div>{"".join(units)}')
    kind = m.get("kind", "")
    parent, _, phref = m.get("parent", "").partition("|")
    crumbs = (f'<a href="{esc(link_target(phref.strip(), ctx))}">{esc(parent.strip())}</a> · {esc(kind)}'
              if parent else f'<a href="{rel(HERE / "index.html", a["dir"])}">Тезисы</a> · сборки')
    hero = (f'<div class="crumbs">{crumbs} · {esc(a["id"])}</div>'
            f'<div class="eyebrow">{"Кейс" if kind == "кейс" else "Карточка сборки · " + esc(kind)}</div>'
            f'<h1>{esc(a["title"])}</h1>'
            + (f'<div class="formula">{inline(m["formula"], ctx)}</div>' if m.get("formula") else "")
            + tags(m, ["frame"]) + tools(a["deck"], a["dir"], a["src"]))
    main = f'<div class="lead doc">{md_html(a["lead"], ctx, skip_quotes=True)}</div>{"".join(body)}'
    return page(f'{a["title"]} · {kind or "сборка"}', hero, main)


def assembly_slides(a, theses, links):
    """Чистые слайды сборки: обложка + слайды тезисов со сквозной нумерацией."""
    ctx = Ctx(a["dir"], a["dir"], links, clean=True)
    m = a["meta"]
    ts = members(a, theses)
    total = 1 + sum(len(t["slides"]) for t in ts)
    items = []
    for _, its in a["groups"]:
        for it in its:
            t = theses.get(it["id"])
            items.append(f'<li>{esc(t["title"])}</li>' if t else
                         f'<li style="opacity:.55">{inline(it["text"], ctx)}</li>')
    h = [f'<section class="slide cover"><div class="s-in">'
         + s_top(f'{m.get("kind", "сборка")} · {a["title"]}', a["out"], ctx, m.get("updated", ""))
         + f'<h2 class="s-title">{esc(a["title"])}</h2><div class="s-body"><ol>{"".join(items)}</ol></div>'
         f'<div class="s-foot"><span>{FOOT}</span><span>1 / {total}</span></div></div></section>']
    n = 1
    for t in ts:
        h.append(slides_seq(t, Ctx(t["dir"], a["dir"], links, clean=True), t["out"], start=n, total=total))
        n += len(t["slides"])
    return deck_page(f'{a["title"]} · слайды', "\n".join(h))


# Каталоги дисциплин — пул единиц-кандидатов в тезисы: файл → (колонка id, название, вид, блок, суть)
UNIT_SOURCES = [
    ("concepts.csv", "concept_id", "title", "рамка", "block", "summary"),
    ("methodologies.csv", "method_id", "title", "методология", "block", "summary"),
    ("business-models.csv", "model_id", "company", "разбор компании", "", "transition_summary"),
    ("cases.csv", "case_id", "title", "прикладной кейс", "block", "summary"),
    ("industries.csv", "industry_id", "industry", "отрасль", "", "post_ai"),
    ("technologies.csv", "tech_id", "name", "технология", "block", "summary"),
    ("topics.csv", "topic_id", "title", "тема → лекция", "block", "summary"),
]
CAT_ID = re.compile(r"`([a-z]+\d+|tech-[\w-]+)`")
BLOCKS = {"transition": "Постинформационный переход", "landscape": "Технологический ландшафт",
          "ai-business": "ИИ и модели бизнеса", "management": "Модели управления"}
# Виды сборок: что это и какие строки каталога (вид единицы) — кандидаты в сборку такого вида
ASSEMBLY_KINDS = [("лекция", "тезисы на одно занятие", ["тема → лекция"]),
                  ("курс", "последовательность лекций", []),
                  ("серия", "тезисы одного вида в одной рамке", []),
                  ("кейс", "разбор компании: страница и слайды", ["разбор компании", "прикладной кейс"]),
                  ("методология", "граф тезисов, разбитый на блоки", ["методология"])]


def catalogs():
    """Папки каталогов дисциплин: [(дисциплина, путь)]."""
    return [(Path(c).parent.name, Path(c)) for c in sorted(glob.glob(str(REPO / "disciplines" / "*" / "catalog")))]


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def units(theses, assemblies):
    """Реестр единиц: собранные тезисы, затем строки каталогов дисциплин, ещё не ставшие тезисами.

    Строка каталога, которую раскрывает тезис (поле covers), вливается в строку тезиса. Остальные строки —
    «в сборке» (запланированы пунктом сборки) или «кандидат» (пока только строка из презентаций).
    """
    covered = {c for t in theses.values() for c in split(t["meta"].get("covers"))}
    planned = {}
    for a in assemblies:
        for _, its in a["groups"]:
            for i in its:
                if i["id"] not in theses:
                    for c in CAT_ID.findall(i["text"]):
                        planned.setdefault(c, []).append(a["id"])
    rows = []
    for t in theses.values():
        m = t["meta"]
        rows.append({"unit_id": t["id"], "discipline": "", "source": "theses", "title": t["title"],
                     "kind": m.get("kind", ""), "block": ";".join(split(m.get("blocks"))),
                     "summary": m.get("formula", ""), "status": "собран", "theses": t["id"],
                     "covers": ";".join(split(m.get("covers"))), "slides": len(t["slides"]),
                     "assemblies": ";".join(a["id"] for a in assemblies if t in members(a, theses)),
                     "source_ids": ""})
    for disc, cat in catalogs():
        for fname, idc, titlec, kind, blockc, sumc in UNIT_SOURCES:
            if not (cat / fname).exists():
                continue
            for r in read_csv(cat / fname):
                uid = r[idc]
                if uid in covered:
                    continue
                asm = planned.get(uid, [])
                rows.append({"unit_id": uid, "discipline": disc, "source": fname, "title": r[titlec],
                             "kind": kind, "block": r.get(blockc, "") if blockc else "",
                             "summary": r.get(sumc, ""), "status": "в сборке" if asm else "кандидат",
                             "theses": "", "covers": "", "slides": "", "assemblies": ";".join(asm),
                             "source_ids": r.get("source_ids", "")})
    return rows


def dset(v):
    """Значения для data-атрибута фильтра: «|a|b|» — строка подходит, если в ней есть выбранное значение."""
    return "|" + "|".join(split(v)) + "|"


UNITS_SCRIPT = """<script>
(function(){
  var rows=[].slice.call(document.querySelectorAll('#units tbody tr')),q=document.getElementById('u-q'),
      sel=[].slice.call(document.querySelectorAll('.u-f')),cnt=document.getElementById('u-n');
  function run(){var s=q.value.trim().toLowerCase(),n=0;
    rows.forEach(function(r){var ok=(!s||r.textContent.toLowerCase().indexOf(s)>=0)&&
      sel.every(function(x){return !x.value||r.dataset[x.dataset.k].indexOf('|'+x.value+'|')>=0;});
      r.style.display=ok?'':'none';if(ok)n++;});cnt.textContent=n;}
  q.oninput=run;sel.forEach(function(x){x.onchange=run;});run();
  // ссылки «кандидаты» у видов сборок выставляют фильтры таблицы
  [].forEach.call(document.querySelectorAll('[data-filter]'),function(a){a.onclick=function(){
    var f=JSON.parse(a.dataset.filter);q.value='';
    sel.forEach(function(x){x.value=f[x.dataset.k]||'';});run();};});
  // длинная суть строки раскрывается по клику
  [].forEach.call(document.querySelectorAll('#units .clamp'),function(e){e.onclick=function(){e.classList.toggle('open');};});
})();
</script>"""


def units_html(rows, theses, assemblies):
    by_id = {a["id"]: a for a in assemblies}

    def options(key, label, names=None):
        vals = sorted({v for r in rows for v in split(r[key])})
        return (f'<select class="u-f" data-k="{key}"><option value="">{label}: все</option>'
                + "".join(f'<option value="{esc(v)}">{esc((names or {}).get(v, v))}</option>' for v in vals)
                + "</select>")

    trs = []
    for r in rows:
        if r["status"] == "собран":
            t = theses[r["theses"]]
            href = rel(t["out"], HERE)
            uid = f'<a href="{href}"><code>{esc(r["unit_id"])}</code></a>'
            st = (f'<a href="{href}">карточка</a> · <a href="{rel(t["deck"], HERE)}">слайды</a> · {r["slides"]} сл.'
                  + "".join(f'<br>→ <a href="{rel(by_id[x]["out"], HERE)}">{esc(by_id[x]["title"])}</a>'
                            for x in split(r["assemblies"]))
                  + (f'<br>раскрывает <code>{esc(r["covers"].replace(";", ", "))}</code>' if r["covers"] else ""))
        elif r["assemblies"]:
            uid = f'<code>{esc(r["unit_id"])}</code>'
            st = " ".join(f'<a href="{rel(by_id[x]["out"], HERE)}">{esc(by_id[x]["title"])}</a>'
                          for x in split(r["assemblies"]))
        else:
            uid = f'<code>{esc(r["unit_id"])}</code>'
            st = (f'<a href="{GH}blob/main/disciplines/{r["discipline"]}/catalog/{r["source"]}">'
                  f'{esc(r["source"])}</a>')
        blocks_ = ", ".join(BLOCKS.get(b, b) for b in split(r["block"]))
        clamp = ' clamp" title="Показать целиком' if len(r["summary"]) > 170 else ""
        cls = "ok" if r["status"] == "собран" else "plan" if r["status"] == "в сборке" else "cand"
        trs.append(f'<tr data-kind="{esc(dset(r["kind"]))}" data-block="{esc(dset(r["block"]))}" '
                   f'data-status="{esc(dset(r["status"]))}"><td>{uid}</td><td><b>{esc(r["title"])}</b>'
                   f'<div class="u-sum{clamp}">{esc(r["summary"])}</div></td><td>{esc(r["kind"])}</td>'
                   f'<td class="u-bl">{esc(blocks_)}</td><td><span class="u-st u-{cls}">{esc(r["status"])}</span>'
                   f'<div class="u-sum">{st}</div></td></tr>')
    stat = {s: sum(1 for r in rows if r["status"] == s) for s in ("собран", "в сборке", "кандидат")}
    return (f'<div class="doc" id="theses"><span class="anchor" id="all"></span>'
            f'<h2>Тезисы · {stat["собран"]} собрано</h2>'
            f'<p class="lead">Сверху — собранные тезисы. Ниже — кандидаты: строки каталогов дисциплин, '
            f'извлечённые из авторских презентаций (концепты, методологии, модели бизнеса, кейсы, отрасли, '
            f'технологии, темы). Кандидат становится тезисом, когда его собирают в карточку на 1–3 слайда; '
            f'строка каталога тогда вливается в строку тезиса. В сборках запланировано — {stat["в сборке"]}, '
            f'кандидатов — {stat["кандидат"]}. Таблица для Excel — '
            f'<a href="{GH}blob/main/theses/catalog/units.csv"><code>catalog/units.csv</code></a>.</p></div>'
            f'<div class="u-bar"><input id="u-q" type="search" placeholder="Поиск по названию и сути…">'
            f'{options("status", "Статус")}{options("kind", "Вид")}{options("block", "Блок", BLOCKS)}'
            f'<span class="u-cnt">показано <b id="u-n">{len(rows)}</b></span></div>'
            f'<div class="doc"><table id="units"><thead><tr><th>Id</th><th>Тезис / единица</th><th>Вид</th>'
            f'<th>Блок</th><th>Статус</th></tr></thead><tbody>{"".join(trs)}</tbody></table></div>{UNITS_SCRIPT}')


def lead_text(a, limit=230):
    """Первый абзац вводного текста сборки — простым текстом, для карточки."""
    para = next((d for k, d in blocks(a["lead"]) if k == "p"), "")
    t = re.sub(r"\s+", " ", plain(para))
    return t if len(t) <= limit else t[:limit - 5].rsplit(" ", 1)[0] + "…"


def assemblies_html(assemblies, theses, rows):
    """Сборки карточками по видам; над ними — виды сборок со счётчиками готовых и кандидатов."""
    order = [k for k, _, _ in ASSEMBLY_KINDS]
    kinds = []
    for kind, what, cand in ASSEMBLY_KINDS:
        n = sum(1 for a in assemblies if a["meta"].get("kind") == kind)
        links = []
        for ck in cand:                          # кандидаты — по каждому виду строк каталога, ссылкой на фильтр
            c = sum(1 for r in rows if r["kind"] == ck and r["status"] == "кандидат")
            flt = esc('{"status":"кандидат","kind":"%s"}' % ck)
            links.append(f'<a href="#theses" data-filter="{flt}">кандидатов: {c}</a>'
                         + (f' <small>({esc(ck)})</small>' if len(cand) > 1 else ""))
        kinds.append(f'<div class="kind"><b>{esc(kind)}</b><span>{esc(what)}</span><em>'
                     + (f'{n} готово' if n else 'пока нет') + "".join(f'<br>{x}' for x in links) + "</em></div>")

    def rank(a):
        k = a["meta"].get("kind", "")
        return (order.index(k) if k in order else len(order), a["title"])

    cards = []
    for a in sorted(assemblies, key=rank):
        m = a["meta"]
        ts = members(a, theses)
        plan = sum(1 for _, its in a["groups"] for i in its if i["id"] not in theses)
        href = rel(a["out"], HERE)
        cards.append(f'<div class="card"><span class="k">{esc(m.get("kind", "сборка"))} · {esc(m.get("status", ""))}</span>'
                     f'<h3><a href="{href}">{esc(a["title"])}</a></h3><span class="fx">{esc(lead_text(a))}</span>'
                     f'<span class="meta">{len(ts)} тез. · {1 + sum(len(t["slides"]) for t in ts)} сл.'
                     + (f' · {plan} планируется' if plan else "") + f' · обновлено {esc(m.get("updated", ""))}</span>'
                     f'<span class="meta"><a href="{href}">карточка</a> · <a href="{rel(a["deck"], HERE)}">слайды</a>'
                     f'</span></div>')
    return (f'<div class="doc" id="assemblies"><h2>Сборки · {len(assemblies)}</h2>'
            f'<p class="lead">Сборка — упорядоченный список тезисов, контент не копируется. Готовые сборки — '
            f'карточками ниже; кандидаты в сборки — строки таблицы тезисов (темы, методологии, кейсы).</p></div>'
            f'<div class="kinds">{"".join(kinds)}</div><div class="cards">{"".join(cards)}</div>')


def catalog_tables():
    """Банк вопросов и источники из каталогов дисциплин — раскрывающимися таблицами."""
    qs, ss, topics = [], [], {}
    for disc, cat in catalogs():
        if (cat / "topics.csv").exists():
            topics.update({r["topic_id"]: r["title"] for r in read_csv(cat / "topics.csv")})
        if (cat / "questions.csv").exists():
            qs += read_csv(cat / "questions.csv")
        if (cat / "sources.csv").exists():
            ss += read_csv(cat / "sources.csv")
    qrows = "".join(f'<tr><td><code>{esc(r["question_id"])}</code></td><td>{esc(r["question"])}'
                    f'<div class="u-sum">{esc(r["topic_id"])} · {esc(topics.get(r["topic_id"], ""))}</div></td>'
                    f'<td class="u-bl">{esc(BLOCKS.get(r["block"], r["block"]))}</td>'
                    f'<td>{esc(r["type"])}<div class="u-sum">{esc(r["level"])}</div></td></tr>' for r in qs)
    srows = "".join(f'<tr><td><code>{esc(r["source_id"])}</code></td><td><b>{esc(r["title"])}</b>'
                    f'<div class="u-sum">{esc(r["note"])}</div></td><td>{esc(r["date"])}</td>'
                    f'<td>{esc(r["slides"])}</td></tr>' for r in ss)
    return (f'<div class="doc" id="questions"><details class="more"><summary>Вопросы · {len(qs)}</summary>'
            f'<p class="lead">Банк вопросов каталога дисциплины — к теме и блоку, с форматом и уровнем. Вопросы '
            f'тезисов — в их карточках.</p><table><thead><tr><th>Id</th><th>Вопрос</th><th>Блок</th>'
            f'<th>Формат</th></tr></thead><tbody>{qrows}</tbody></table></details>'
            f'<details class="more" id="sources"><summary>Источники · {len(ss)}</summary>'
            f'<p class="lead">Авторские презентации, из которых извлечены кандидаты. Оригиналы приватные и в '
            f'репозиторий не кладутся.</p><table><thead><tr><th>Id</th><th>Презентация</th><th>Дата</th>'
            f'<th>Слайдов</th></tr></thead><tbody>{srows}</tbody></table></details></div>')


def index_page(theses, assemblies, rows):
    hero = ('<div class="crumbs"><a href="../../index.html">Мастерская</a> · тезисы и сборки</div>'
            '<div class="eyebrow">База контента мастерской</div><h1>Тезисы и сборки</h1>'
            '<div class="formula">Тезис — базовая единица контента: одна мысль, одна формула, 1–3 слайда. '
            'Из тезисов собираются лекции, курсы, серии, кейсы и методологии.</div>'
            f'<div class="tools"><a class="btn" href="#assemblies">Сборки · {len(assemblies)}</a>'
            f'<a class="btn" href="#theses">Тезисы · {len(theses)}</a>'
            f'<a class="btn ghost" href="#questions">Вопросы и источники</a>'
            f'<a class="btn ghost" href="{GH}blob/main/theses/README.md">Модель слоя</a></div>')
    main = assemblies_html(assemblies, theses, rows) + units_html(rows, theses, assemblies) + catalog_tables()
    return page("Тезисы и сборки · мастерская", hero, main)


# ---------------------------------------------------------------- индексы CSV

def write_csv(name, header, rows):
    (ROOT / "catalog").mkdir(exist_ok=True)
    with open(ROOT / "catalog" / name, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


# ---------------------------------------------------------------- экспорт

def browser():
    cands = [os.environ.get("THESES_BROWSER", "")]
    cands += [shutil.which(n) or "" for n in ("chrome", "google-chrome", "chromium", "chromium-browser", "msedge")]
    cands += [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
              "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
              "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"]
    for c in cands:
        if c and os.path.exists(c):
            return c
    sys.exit("Не найден Chrome/Edge для экспорта. Укажите путь в переменной THESES_BROWSER.")


def headless(args, url, out):
    """Запускает браузер и ждёт файл: на Windows msedge.exe возвращается раньше, чем допишет результат."""
    out.unlink(missing_ok=True)
    prof = tempfile.mkdtemp(prefix="theses-browser-")
    subprocess.run([browser(), "--headless=new", "--disable-gpu", "--hide-scrollbars",
                    f"--user-data-dir={prof}", "--virtual-time-budget=8000", *args, url],
                   check=True, capture_output=True)
    size, deadline = -1, time.time() + 90
    while time.time() < deadline:
        cur = out.stat().st_size if out.exists() else -1
        if cur > 0 and cur == size:
            break
        size = cur
        time.sleep(0.7)
    else:
        sys.exit(f"Браузер не создал {out}")
    shutil.rmtree(prof, ignore_errors=True)


def export_pdf(deck, name):
    out = EXPORT / f"{name}.pdf"
    headless(["--no-pdf-header-footer", f"--print-to-pdf={out}"], deck.as_uri() + "?export", out)
    print("  pdf ", out.relative_to(REPO))


def link_box(prs, sl, url, label):
    """Невидимая кликабельная область поверх индекса на кадре слайда → карточка на GitHub Pages."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    px = prs.slide_width / 1280
    w = min(len(label) * 10 + 24, 900)
    shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(44 * px), int(28 * px), int(w * px), int(34 * px))
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(255, 255, 255)
    clr = shp._element.spPr.find(qn("a:solidFill"))[0]
    clr.append(clr.makeelement(qn("a:alpha"), {"val": "0"}))
    shp.line.fill.background()
    shp.click_action.hyperlink.address = url
    shp.name = "Ссылка на карточку"


def export_pptx(deck, name, frames_info):
    """frames_info: [(заметки, url карточки, подпись индекса)] по слайдам."""
    from pptx import Presentation
    from pptx.util import Inches
    frames = EXPORT / "frames" / name
    frames.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for k, (notes, url, label) in enumerate(frames_info):
        png = frames / f"{k + 1:02d}.png"
        headless(["--window-size=1280,720", "--force-device-scale-factor=2", f"--screenshot={png}"],
                 f"{deck.as_uri()}?frame={k + 1}", png)
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        link_box(prs, sl, url, label)
        sl.notes_slide.notes_text_frame.text = notes
    out = EXPORT / f"{name}.pptx"
    prs.save(out)
    print("  pptx", out.relative_to(REPO))


def frames_info(t):
    label = f'{t["id"]} · {t["title"]}'
    return [(f'{s["title"]}\n\n{plain(s["md"])}' + (f'\n\n— Комментарий —\n{plain(s["notes"])}' if s["notes"] else ""),
             pages_url(t["out"]), label) for s in t["slides"]]


# ---------------------------------------------------------------- main

def main():
    want_pdf = "--pdf" in sys.argv or "--export" in sys.argv
    want_pptx = "--pptx" in sys.argv or "--export" in sys.argv

    theses = {}
    for readme in sorted(glob.glob(str(ROOT / "th*" / "README.md"))):
        t = load_thesis(readme)
        if t["id"] in theses:
            warn(f'{t["id"]}: повторяющийся id')
        theses[t["id"]] = t
    assemblies = [load_assembly(p) for p in sorted(glob.glob(str(ROOT / "assemblies" / "*.md"))) + case_sources()]

    links = {t["src"]: t["out"] for t in theses.values()}
    links.update({a["src"]: a["out"] for a in assemblies})
    member = {tid: [a for a in assemblies if any(i["id"] == tid for _, its in a["groups"] for i in its)]
              for tid in theses}

    def write(path, text):
        path.write_text(text, encoding="utf-8")
        print("html", path.relative_to(REPO))

    for t in theses.values():
        write(t["out"], thesis_page(t, links, member[t["id"]]))
        write(t["deck"], thesis_slides(t, links))
    for a in assemblies:
        write(a["out"], assembly_page(a, theses, links))
        write(a["deck"], assembly_slides(a, theses, links))
    rows = units(theses, assemblies)
    write(HERE / "index.html", index_page(theses, assemblies, rows))
    write_csv("units.csv", list(rows[0].keys()), [list(r.values()) for r in rows])

    write_csv("theses.csv",
              ["thesis_id", "title", "formula", "kind", "blocks", "refs", "questions", "assignments", "dossier",
               "slides", "assemblies", "status", "version", "updated", "file"],
              [[t["id"], t["title"], t["meta"].get("formula", ""), t["meta"].get("kind", ""),
                t["meta"].get("blocks", ""), t["meta"].get("refs", ""), t["meta"].get("questions", ""),
                t["meta"].get("assignments", ""), t["meta"].get("dossier", ""), len(t["slides"]),
                ";".join(a["id"] for a in member[t["id"]]), t["meta"].get("status", ""),
                t["meta"].get("version", ""), t["meta"].get("updated", ""),
                t["src"].relative_to(ROOT).as_posix()] for t in theses.values()])
    write_csv("assemblies.csv",
              ["assembly_id", "title", "kind", "frame", "theses", "planned", "status", "updated", "file"],
              [[a["id"], a["title"], a["meta"].get("kind", ""), a["meta"].get("frame", ""),
                ";".join(t["id"] for t in members(a, theses)),
                ";".join(plain(i["text"]) for _, its in a["groups"] for i in its if i["id"] not in theses),
                a["meta"].get("status", ""), a["meta"].get("updated", ""),
                rel(a["src"], ROOT)] for a in assemblies])
    print("csv  theses/catalog/theses.csv, assemblies.csv, units.csv")

    if want_pdf or want_pptx:
        EXPORT.mkdir(exist_ok=True)
        for t in theses.values():
            name = t["dir"].name
            if want_pdf:
                export_pdf(t["deck"], name)
            if want_pptx:
                export_pptx(t["deck"], name, frames_info(t))
        for a in assemblies:
            if want_pdf:
                export_pdf(a["deck"], a["id"])
            if want_pptx:
                cover = (a["title"], pages_url(a["out"]), f'{a["meta"].get("kind", "сборка")} · {a["title"]}')
                export_pptx(a["deck"], a["id"], [cover] + [f for t in members(a, theses) for f in frames_info(t)])

    for w in WARN:
        print("ВНИМАНИЕ:", w)


if __name__ == "__main__":
    main()
